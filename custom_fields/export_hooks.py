"""
Shared helpers so custom fields appear in Horilla's Select Columns to Export
modal and are written into exported files.

Horilla's export catalog only knows about model columns. The functions here
are consumed by extension registrations, not applied as monkey-patches
themselves:

- ``get_available_models``/``export_model_data`` are overridden by
  ``CustomFieldExportViewExtension`` in ``custom_fields/view_extensions.py``
  (``ExportView`` is dispatched via ``as_view()``, so ``ViewExtension``/
  ``_inherit_view`` applies).
- ``_install_export_properties``/``_uninstall_export_properties`` are used
  by ``CustomFieldBulkExportExtension`` in
  ``custom_fields/mixin_extensions.py`` (``HorillaBulkExportMixin`` is a
  bare mixin, and ``get_export_cell_value`` a bare module function — neither
  is ever dispatched or resolved, so both are extended through
  ``MixinExtension``/``_inherit_mixin`` instead). They install/uninstall
  ``cf_*`` properties on the target model only for the duration of one
  ``handle_export`` call, paired in a ``try``/``finally`` — scoped to that
  one model class, never Django internals or another app's shared state.
"""

import csv
import logging
from io import BytesIO, StringIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from custom_fields.detail_hooks import custom_field_selector_items
from horilla.contrib.core.utils import sanitize_export_value

logger = logging.getLogger(__name__)


class _MaterializedObjectList(list):
    """
    A plain ``list`` that also answers ``.iterator(chunk_size=...)``.

    ``export_model_data`` on ``ExportView`` calls ``queryset.iterator(...)``
    to stream rows without caching the full result set. Custom-field export
    must pass a plain, already-materialized list of objects instead of a
    real ``QuerySet`` (the objects already carry ``cf_*`` values attached in
    memory by ``attach_custom_field_values_to_objects``, so they cannot be
    re-queried) — this thin subclass keeps that in-memory list working with
    the streaming call site.
    """

    def iterator(self, chunk_size=None):
        """Yield materialized objects without re-querying the database."""
        return iter(self)


def _cf_property(name):
    return property(lambda obj, key=name: obj.__dict__.get(key, ""))


def _install_export_properties(model, extras):
    """Expose ``cf_*`` names as properties so Horilla's export catalog includes them."""
    installed = []
    for _label, name in extras:
        if not hasattr(model, name):
            setattr(model, name, _cf_property(name))
            installed.append(name)
    old_labels = getattr(model, "PROPERTY_LABELS", None)
    labels = dict(old_labels or {})
    for label, name in extras:
        labels[name] = label
    model.PROPERTY_LABELS = labels
    return installed, old_labels


def _uninstall_export_properties(model, installed, old_labels):
    for name in installed:
        if hasattr(model, name):
            delattr(model, name)
    if old_labels is None:
        if hasattr(model, "PROPERTY_LABELS"):
            delattr(model, "PROPERTY_LABELS")
    else:
        model.PROPERTY_LABELS = old_labels


def _cell_values(obj, extra_pairs):
    values = []
    for _label, name in extra_pairs:
        raw = obj.__dict__.get(name, "")
        values.append(sanitize_export_value("" if raw is None else str(raw)))
    return values


def add_custom_fields_to_export_modules(modules):
    """Add custom fields to Settings → Export Data column pickers."""
    from horilla.apps import apps

    for module in modules or []:
        app_label = module.get("app_label")
        model_name = module.get("name")
        if not app_label or not model_name:
            continue
        try:
            model = apps.get_model(app_label, model_name)
        except LookupError:
            continue
        extras = custom_field_selector_items(model)
        if not extras:
            continue
        fields = list(module.get("fields") or [])
        existing = {item.get("name") for item in fields}
        for label, name in extras:
            if name not in existing:
                fields.append({"name": name, "label": label})
                existing.add(name)
        module["fields"] = fields
    return modules


def _custom_fields_only_export(view, model, export_format, objects, extra_pairs):
    """Build a csv/xlsx buffer that contains only custom-field columns."""
    headers = [str(label) for label, _name in extra_pairs]
    rows = [_cell_values(obj, extra_pairs) for obj in objects]
    filename = view.get_export_filename(model, export_format)
    if export_format == "csv":
        buffer = StringIO()
        writer = csv.writer(buffer)
        writer.writerow(headers)
        writer.writerows(rows)
        return filename, BytesIO(buffer.getvalue().encode("utf-8"))
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    header_font = Font(bold=True)
    header_alignment = Alignment(horizontal="center")
    header_fill = PatternFill(
        start_color="eafb5b", end_color="eafb5b", fill_type="solid"
    )
    for cell in sheet[1]:
        cell.font = header_font
        cell.alignment = header_alignment
        cell.fill = header_fill
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return filename, buffer


def _append_custom_field_columns(data, export_format, objects, extra_pairs):
    """Append custom-field columns onto an already-built csv/xlsx export buffer."""
    extra_headers = [str(label) for label, _name in extra_pairs]
    extra_values = [_cell_values(obj, extra_pairs) for obj in objects]
    raw = data.getvalue()
    if export_format == "csv":
        text = raw.decode("utf-8") if isinstance(raw, (bytes, bytearray)) else raw
        reader = csv.reader(StringIO(text))
        rows = list(reader)
        if not rows:
            rows = [[]]
        rows[0] = list(rows[0]) + extra_headers
        for index, extra_row in enumerate(extra_values):
            if index + 1 < len(rows):
                rows[index + 1] = list(rows[index + 1]) + extra_row
            else:
                rows.append(extra_row)
        buffer = StringIO()
        csv.writer(buffer).writerows(rows)
        return BytesIO(buffer.getvalue().encode("utf-8"))

    workbook = load_workbook(BytesIO(raw))
    sheet = workbook.active
    start_col = sheet.max_column + 1
    for offset, header in enumerate(extra_headers):
        sheet.cell(row=1, column=start_col + offset, value=header)
    for row_index, extra_row in enumerate(extra_values, start=2):
        for offset, value in enumerate(extra_row):
            sheet.cell(row=row_index, column=start_col + offset, value=value)
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer
